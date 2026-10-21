import io
import os
import re
import sys
import tempfile
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from openpyxl import load_workbook

import db
import exporter
from content import CATEGORIES, CONTENT, SIGNS
from engine import _redundant, generate_month, split_sentences


def _sentences(entries):
    return [s for e in entries for s in split_sentences(e.text)]


def test_no_sentence_repeats_across_months_and_history():
    p = os.path.join(tempfile.mkdtemp(), "t.db")
    seen = set()
    for m in range(1, 13):  # un año completo de historial
        es = generate_month(2026, m, db.used_sentences(p))
        sents = _sentences(es)
        assert len(sents) == len(set(sents)), "oración repetida dentro del mes"
        assert not seen & set(sents), "oración repetida respecto al historial"
        seen |= set(sents)
        db.save_horoscope(2026, m, es, p)
        assert len({e.text for e in es}) == len(es)


def test_key_ideas_not_repeated_within_sign_and_month():
    es = generate_month(2026, 10, set())
    for sign in SIGNS:
        for cat in CATEGORIES:
            c = CONTENT[cat]
            used = Counter()
            for e in es:
                if e.sign != sign or e.category != cat:
                    continue
                low = e.text.lower()
                used.update(("sit", p) for p in c["sit"] if p in low)
                used.update(("eff", p) for p in c["eff"] if p in low)
                used.update(("light", p) for p in c["light"] if p in low)
            assert used and max(used.values()) <= 1, (sign, cat, used.most_common(1))


def test_no_redundancy_inside_a_text():
    es = generate_month(2026, 10, set())
    for e in es:
        parts = split_sentences(e.text)
        assert 1 <= len(parts) <= 3
        assert not _redundant(parts), e.text
        assert len(re.findall(r"\bpero\b", e.text.lower())) <= 1, e.text


def test_versions_and_delete():
    p = os.path.join(tempfile.mkdtemp(), "t.db")
    es = generate_month(2026, 12, set())
    db.save_horoscope(2026, 12, es, p)
    h2 = db.save_horoscope(2026, 12, es[:5], p)
    assert sorted(h["version"] for h in db.list_horoscopes(p)) == [1, 2]
    db.delete_horoscope(h2, p)
    assert len(db.list_horoscopes(p)) == 1


def test_excel_format():
    es = generate_month(2026, 10, set())
    rows = [dict(sign=e.sign, category=e.category, day=e.day.isoformat(), text=e.text) for e in es]
    ws = load_workbook(io.BytesIO(exporter.build_xlsx(rows))).active
    assert ws.title == "HORO" and ws.max_row == 1489
    assert [c.value for c in ws[1]] == ["SIGNO", "CATEGORIA", "PRONOSTICO", "FECHA"]
    assert ws["A2"].value == "ARIES" and ws["B2"].value == "Hoy"


def _rows(entries):
    return [dict(sign=e.sign, category=e.category, day=e.day.isoformat(), text=e.text) for e in entries]


def test_review_passes_clean_month_and_returns_report():
    from review import generate_reviewed
    entries, data, rep = generate_reviewed(2026, 11, set(), seed=1)
    assert len(entries) == 1440 and data and rep["attempts"] >= 1
    assert [x["pass"] for x in rep["passes"]] == [1, 2] and all(x["count"] == 0 for x in rep["passes"])


def test_audit_detects_injected_errors():
    from review import audit_rows
    es = generate_month(2026, 11, set(), seed=2)
    rows = _rows(es)
    assert audit_rows(rows, 2026, 11) == []
    # fila faltante
    assert any("Falta" in p or "Cantidad" in p for p in audit_rows(rows[:-1], 2026, 11))
    # texto duplicado
    dup = [dict(r) for r in rows]
    dup[1]["text"] = dup[0]["text"]
    assert any("repetid" in p for p in audit_rows(dup, 2026, 11))
    # oración ya usada en el historial
    hist = set(split_sentences(rows[0]["text"]))
    assert any("anteriores" in p for p in audit_rows(rows, 2026, 11, hist))
    # defecto de formato
    bad = [dict(r) for r in rows]
    bad[5]["text"] = bad[5]["text"].replace(" ", "  ", 1)
    assert any("Formato" in p for p in audit_rows(bad, 2026, 11))


def test_review_redoes_month_when_first_attempt_is_defective(monkeypatch):
    import review
    from engine import Entry
    real = review.generate_month
    calls = []

    def flaky(year, month, history, seed=None):
        entries = real(year, month, history, seed=seed)
        calls.append(1)
        if len(calls) == 1:  # primer intento defectuoso: texto duplicado
            e0, e1 = entries[0], entries[1]
            entries[1] = Entry(e1.sign, e1.category, e0.text, e1.day)
        return entries

    monkeypatch.setattr(review, "generate_month", flaky)
    entries, data, rep = review.generate_reviewed(2026, 11, set(), seed=3)
    assert rep["attempts"] == 2 and len(calls) == 2
    assert rep["log"][0]["count"] > 0 and rep["passes"][-1]["count"] == 0
