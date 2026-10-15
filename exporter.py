"""Exporta un horóscopo al mismo formato del Excel de referencia (hoja HORO)."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook

from engine import MONTH_ABBR


def build_xlsx(entries):
    """entries: iterable de dicts con sign, category, day (iso), text. Devuelve bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "HORO"
    ws.append(["SIGNO", "CATEGORIA", "PRONOSTICO", "FECHA"])
    for e in entries:
        ws.append([e["sign"], e["category"], e["text"], date.fromisoformat(e["day"])])
        ws.cell(row=ws.max_row, column=4).number_format = "mm-dd-yy"
    ws.column_dimensions["C"].width = 106.5546875
    ws.column_dimensions["D"].width = 23.6640625
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(year, month, version):
    return f"Horoscopo_{MONTH_ABBR[month - 1]}_{year}-V{version}.xlsx"
