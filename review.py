"""Revisión automática de un horóscopo antes de entregarlo.

Dos pasadas independientes del generador:
  1. Revisa los textos recién generados.
  2. Vuelve a revisar desde cero lo que realmente se entregará: el Excel ya armado,
     leído de vuelta, contra el historial guardado.
Si alguna pasada encuentra problemas, el mes se rehace (hasta MAX_ATTEMPTS veces) y
nada se guarda ni se ofrece para descargar hasta que ambas pasadas salgan limpias.
"""
import calendar
import io
import re
from collections import Counter
from datetime import date

from openpyxl import load_workbook

from content import CATEGORIES, CONTENT, SIGNS
from engine import _redundant, generate_month, split_sentences
from exporter import build_xlsx

MAX_ATTEMPTS = 12
MAX_REPORTED = 8  # problemas que se detallan por pasada


class ReviewError(RuntimeError):
    pass


def audit_rows(rows, year, month, history=()):
    """rows: iterable de dicts sign, category, day, text. Devuelve la lista de problemas."""
    rows = list(rows)
    history = set(history)
    problems = []
    n_days = calendar.monthrange(year, month)[1]

    # 1. Cobertura: todos los signos x días x categorías, sin faltantes ni sobrantes.
    expected = {(s, c, date(year, month, d).isoformat()) for s in SIGNS for c in CATEGORIES for d in range(1, n_days + 1)}
    got = Counter((r["sign"], r["category"], str(r["day"])[:10]) for r in rows)
    if len(rows) != len(expected):
        problems.append(f"Cantidad de filas {len(rows)} (se esperaban {len(expected)})")
    for key in sorted(expected - set(got))[:3]:
        problems.append(f"Falta el pronóstico {key}")
    for key, n in got.items():
        if key not in expected:
            problems.append(f"Fila fuera del mes o con signo/categoría inválidos: {key}")
        elif n > 1:
            problems.append(f"Pronóstico duplicado para {key}")

    # 2. Forma de cada texto.
    for r in rows:
        t, where = r["text"], f"{r['sign']} {r['category']} {str(r['day'])[:10]}"
        parts = split_sentences(t or "")
        if not t or not t.strip():
            problems.append(f"Texto vacío en {where}")
            continue
        if "  " in t or "{" in t or "}" in t or "None" in t or t != t.strip():
            problems.append(f"Formato defectuoso en {where}")
        if not 1 <= len(parts) <= 3:
            problems.append(f"{len(parts)} oraciones en {where}")
        if any(not p[0].isupper() or not p.endswith(".") for p in parts):
            problems.append(f"Oración sin mayúscula inicial o sin punto final en {where}")
        if _redundant(parts):
            problems.append(f"Ideas o palabras repetidas dentro del texto en {where}")

    # 3. Ninguna repetición: textos y oraciones únicos en el mes y frente al historial.
    texts = Counter(r["text"] for r in rows)
    sents = Counter(s for r in rows for s in split_sentences(r["text"] or ""))
    for t, n in texts.items():
        if n > 1:
            problems.append(f"Texto repetido {n} veces: «{t[:60]}…»")
    for s, n in sents.items():
        if n > 1:
            problems.append(f"Oración repetida {n} veces en el mes: «{s[:70]}»")
        if s in history:
            problems.append(f"Oración ya usada en versiones anteriores: «{s[:70]}»")

    # 4. La idea central (situación, esfuerzo, resultado) no se repite dentro de un signo/categoría.
    by = {}
    for r in rows:
        by.setdefault((r["sign"], r["category"]), []).append(split_sentences(r["text"] or ""))
    for (sign, cat), texts_ in by.items():
        c = CONTENT[cat]
        cnt = Counter()
        for parts in texts_:
            low = " ".join(parts).lower()
            cnt.update(("situación", p) for p in c["sit"] if p in low)
            cnt.update(("esfuerzo", p) for p in c["eff"] if p in low)
            cnt.update(("resultado", p) for p in c["light"] if p in low)
        for (kind, phrase), n in cnt.items():
            if n > 1:
                problems.append(f"{kind.capitalize()} repetida {n} veces en {sign}/{cat}: «{phrase}»")
    return problems


def _rows_from_entries(entries):
    return [dict(sign=e.sign, category=e.category, day=e.day.isoformat(), text=e.text) for e in entries]


def _rows_from_xlsx(data):
    ws = load_workbook(io.BytesIO(data)).active
    if [c.value for c in ws[1]] != ["SIGNO", "CATEGORIA", "PRONOSTICO", "FECHA"]:
        return None
    return [dict(sign=a, category=b, text=c, day=d) for a, b, c, d in ws.iter_rows(min_row=2, values_only=True)]


def generate_reviewed(year, month, history, seed=None, progress=None):
    """Genera el mes y lo entrega solo tras dos revisiones limpias.

    Devuelve (entries, xlsx_bytes, report). report = {"attempts": n, "passes": [...], "log": [...]}.
    progress: función opcional progress(fracción 0-1, texto) para mostrar avance.
    """
    log = []

    def step(frac, text):
        if progress:
            progress(frac, text)

    for attempt in range(1, MAX_ATTEMPTS + 1):
        step(0.05, "Escribiendo los pronósticos…" if attempt == 1 else f"Corrigiendo y reintentando ({attempt})…")
        entries = generate_month(year, month, history, seed=None if seed is None else seed + attempt)
        step(0.45, "Revisión 1 de 2…")
        p1 = audit_rows(_rows_from_entries(entries), year, month, history)
        log.append({"attempt": attempt, "pass": 1, "problems": p1[:MAX_REPORTED], "count": len(p1)})
        if p1:
            continue
        step(0.7, "Armando el Excel…")
        data = build_xlsx(_rows_from_entries(entries))
        step(0.85, "Revisión 2 de 2…")
        rows2 = _rows_from_xlsx(data)
        p2 = ["El Excel generado no tiene el formato esperado"] if rows2 is None else audit_rows(rows2, year, month, history)
        log.append({"attempt": attempt, "pass": 2, "problems": p2[:MAX_REPORTED], "count": len(p2)})
        if p2:
            continue
        step(1.0, "Listo")
        return entries, data, {"attempts": attempt, "passes": [x for x in log if x["attempt"] == attempt], "log": log}
    detail = "; ".join(p for x in log[-2:] for p in x["problems"][:3])
    raise ReviewError(f"No se pudo entregar un horóscopo limpio tras {MAX_ATTEMPTS} intentos. {detail}")
